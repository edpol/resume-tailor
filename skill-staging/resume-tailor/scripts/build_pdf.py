#!/usr/bin/env python3
"""
build_pdf.py — Swap rewritten highlights into the Word resume template and export as PDF.

Uses the original .docx file as a formatting template, replaces the HIGHLIGHTS bullets
in-place while preserving all Word styles, fonts, and spacing, then converts to PDF
via LibreOffice headless.

Usage:
    python build_pdf.py \
        --template path/to/"Edward Pol Resume.docx" \
        --new-bullets "Rewritten bullet 1" "Rewritten bullet 2" ... \
        --output path/to/output.pdf

    # For backward compatibility, --resume and --old-bullets are also accepted but ignored
    # (the new bullets always replace the 5 HIGHLIGHTS paragraphs in order)

Requires:
    pip install python-docx
    libreoffice (for PDF conversion)
"""

import argparse
import subprocess
import sys
import os
import shutil
import tempfile
from pathlib import Path

# python-docx
try:
    from docx import Document
    from docx.oxml.ns import qn
    from docx.shared import Pt, Inches
    from lxml import etree
except ImportError:
    print("Error: python-docx not installed. Run: pip install python-docx")
    sys.exit(1)


# The 5 HIGHLIGHTS bullet paragraphs are at these fixed indices in the Word doc
# (verified by inspecting paragraph structure — index 2 is "HIGHLIGHTS" header,
#  3-7 are the five bullets, 8 is "PROFESSIONAL EXPERIENCE")
HIGHLIGHTS_PARA_INDICES = [3, 4, 5, 6, 7]


def clear_and_set_paragraph_text(para, new_text: str):
    """
    Replace all runs in a paragraph with a single run containing new_text,
    preserving paragraph-level formatting (indent, spacing, style).
    Copies run-level formatting from the first existing run.
    """
    # Capture formatting from the first run before clearing
    first_run = para.runs[0] if para.runs else None
    font_name = None
    font_size = None
    bold = None
    italic = None

    if first_run:
        font_name = first_run.font.name
        font_size = first_run.font.size
        bold = first_run.bold
        italic = first_run.italic

    # Remove all existing runs by clearing <w:r> elements from the paragraph XML
    p_elem = para._p
    for r in p_elem.findall(qn('w:r')):
        p_elem.remove(r)
    # Also remove hyperlinks and other inline content (keep paragraph properties)
    for child in list(p_elem):
        tag = child.tag
        if tag not in (qn('w:pPr'),):  # keep paragraph properties
            p_elem.remove(child)

    # Add a new run with the replacement text
    run = para.add_run(new_text)
    if font_name:
        run.font.name = font_name
    if font_size:
        run.font.size = font_size
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def swap_highlights(template_path: str, new_bullets: list[str], output_docx: str):
    """Open the Word template, replace the 5 HIGHLIGHTS bullets, save to output_docx."""
    doc = Document(template_path)
    paras = doc.paragraphs

    if len(new_bullets) != 5:
        raise ValueError(f"Expected exactly 5 new bullets, got {len(new_bullets)}")

    for bullet_idx, para_idx in enumerate(HIGHLIGHTS_PARA_INDICES):
        if para_idx >= len(paras):
            raise IndexError(f"Paragraph index {para_idx} out of range (doc has {len(paras)} paragraphs)")
        para = paras[para_idx]
        clear_and_set_paragraph_text(para, new_bullets[bullet_idx])

    doc.save(output_docx)
    print(f"Tailored .docx saved: {output_docx}")


def convert_docx_to_pdf(docx_path: str, output_pdf: str):
    """Convert .docx to PDF using LibreOffice headless."""
    output_dir = str(Path(output_pdf).parent)
    cmd = [
        "libreoffice", "--headless", "--convert-to", "pdf",
        "--outdir", output_dir,
        docx_path
    ]
    result = subprocess.run(cmd, capture_output=True, text=True)
    if result.returncode != 0:
        raise RuntimeError(f"LibreOffice conversion failed:\n{result.stderr}")

    # LibreOffice names the output after the input file
    stem = Path(docx_path).stem
    libreoffice_output = Path(output_dir) / f"{stem}.pdf"
    if not libreoffice_output.exists():
        raise FileNotFoundError(f"Expected LibreOffice output not found: {libreoffice_output}")

    # Rename to the requested output path
    if str(libreoffice_output) != output_pdf:
        shutil.move(str(libreoffice_output), output_pdf)

    print(f"PDF saved: {output_pdf}")


def main():
    parser = argparse.ArgumentParser(
        description="Build a tailored resume PDF from the Word template"
    )
    # Primary interface
    parser.add_argument("--template", help="Path to the Word resume template (.docx)")
    parser.add_argument("--new-bullets", nargs="+", required=True,
                        help="5 replacement highlight bullets (in order)")
    parser.add_argument("--output", required=True, help="Output PDF path")

    # Legacy compatibility (ignored but accepted so old callers don't break)
    parser.add_argument("--resume", help="(legacy) Path to master_resume.md — ignored")
    parser.add_argument("--old-bullets", nargs="+", default=[],
                        help="(legacy) Original bullet text — ignored")

    args = parser.parse_args()

    # Resolve template path
    template = args.template
    if not template:
        # Default: look for the .docx in the assets folder relative to this script
        script_dir = Path(__file__).parent
        template = str(script_dir.parent / "assets" / "Edward Pol Resume.docx")

    if not Path(template).exists():
        print(f"Error: template file not found: {template}")
        sys.exit(1)

    if len(args.new_bullets) != 5:
        print(f"Error: exactly 5 --new-bullets required, got {len(args.new_bullets)}")
        sys.exit(1)

    # Work in a temp dir so we don't pollute the assets folder
    with tempfile.TemporaryDirectory() as tmp_dir:
        tmp_docx = str(Path(tmp_dir) / "resume_tailored.docx")
        print(f"Opening template: {template}")
        swap_highlights(template, args.new_bullets, tmp_docx)
        print(f"Converting to PDF...")
        convert_docx_to_pdf(tmp_docx, args.output)

    print(f"Done! PDF saved to: {args.output}")


if __name__ == "__main__":
    main()
